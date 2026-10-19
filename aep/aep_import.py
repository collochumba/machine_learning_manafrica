"""Template import for the Audit Engagement Program: parse pasted/uploaded tables, map columns, build sections.
Pure functions (no Streamlit) so they can be tested. Nothing here writes to the database."""
import csv
import io
import re

FIELDS = ["area", "requirements", "procedure", "control"]
LABELS = {"area": "Audit Area", "requirements": "Requirements", "procedure": "Procedure / Test", "control": "Control",
          "section": "Main Audit Area heading (optional)"}
ALIASES = {
    "area": ["audit area", "audit objective", "audit objectives", "objective", "area", "process", "subject",
             "audit objective / audit area", "audit area / objective", "audit area/objective"],
    "requirements": ["requirements", "requirement", "criteria", "criterion", "audit criteria", "standard", "standards",
                     "policy requirement", "legal requirement", "requirements / criteria"],
    "procedure": ["procedure / test", "procedure/test", "audit procedure", "audit procedures", "test procedure", "test procedures",
                  "tests to perform", "test to perform", "procedure", "procedures", "test", "tests", "audit test", "audit tests",
                  "audit steps", "steps"],
    "control": ["control", "controls", "existing control", "existing controls", "control activity", "control activities",
                "control assessed", "controls to assess", "key control", "key controls"],
    "section": ["main audit area", "section", "heading", "category", "main area", "audit section"],
}
IGNORE = {"no", "no.", "#", "s/n", "sn", "s.no", "s.no.", "item", "ref", "ref.", "sr", "sr."}


class ImportErr(ValueError):
    pass


def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower()).strip(" :")


def _clean(rows):
    rows = [[("" if c is None else str(c).strip()) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        raise ImportErr("No table content was found. Paste at least a header row and one data row.")
    w = max(len(r) for r in rows)
    return [r + [""] * (w - len(r)) for r in rows]


def parse_text(text):
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip("\n")
    if not text.strip():
        raise ImportErr("Nothing was pasted.")
    if "\t" in text:
        rows = list(csv.reader(io.StringIO(text), delimiter="\t"))
    elif text.count("|") >= 2:  # markdown-style table
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in text.split("\n")
                if not re.fullmatch(r"[\s|:\-]+", ln)]
    else:
        try:
            d = csv.Sniffer().sniff(text[:2000], delimiters=",;")
        except csv.Error:
            raise ImportErr("The pasted text does not look like a table. Copy the table cells from Excel or Word "
                            "(so columns are tab-separated) and paste again.") from None
        rows = list(csv.reader(io.StringIO(text), d))
    rows = _clean(rows)
    if len(rows[0]) < 2:
        raise ImportErr("Only one column was detected. Copy the whole table, including every column, and paste again.")
    return rows


def parse_csv(data):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return parse_text(data.decode(enc).replace("\t", "\t"))
        except UnicodeDecodeError:
            continue
    raise ImportErr("The CSV file encoding could not be read.")


def xlsx_sheets(data):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportErr("This Excel file could not be opened. Save it as .xlsx and try again, or paste the table.") from None
    return wb.sheetnames


def parse_xlsx(data, sheet):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        ws = wb[sheet]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # merged cells: carry the top-left value across the merged range so section headings are not lost
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if r - 1 < len(rows) and c - 1 < len(rows[r - 1]):
                        rows[r - 1][c - 1] = v
    except ImportErr:
        raise
    except Exception:
        raise ImportErr("That worksheet could not be read. Paste the table instead.") from None
    return _clean(rows)


def docx_tables(data):
    """Return the tables of a .docx as a list of row-lists (one entry per table with 2+ columns)."""
    import docx
    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception:
        raise ImportErr("This Word file could not be opened. Save it as .docx, or copy the table and paste it.") from None
    out = []
    for t in doc.tables:
        rows = []
        for r in t.rows:
            cells, seen = [], None
            for c in r.cells:
                cells.append(c.text.strip() if c._tc is not seen else c.text.strip())
                seen = c._tc
            rows.append(cells)
        try:
            rows = _clean(rows)
        except ImportErr:
            continue
        if len(rows[0]) >= 2:
            out.append(rows)
    if not out:
        raise ImportErr("No tables were found in this Word file. Copy the table from Word and use the paste box instead.")
    return out


def guess_header(rows):
    """Index of the header row (one of the first 6 rows with 2+ recognised headings), else None."""
    for i, r in enumerate(rows[:6]):
        if len(auto_map(r, strict=True)) >= 2:
            return i
    return None


def auto_map(headers, strict=False):
    """Map field -> column index using exact aliases first, then (unless strict) 'contains' matches."""
    got, used = {}, set()
    hs = [norm(h) for h in headers]
    for f in FIELDS + ["section"]:
        for j, h in enumerate(hs):
            if j not in used and h in ALIASES[f]:
                got[f] = j; used.add(j); break
    if not strict:
        for f in FIELDS + ["section"]:
            if f in got:
                continue
            for j, h in enumerate(hs):
                if j not in used and h and any(a in h or (len(h) > 3 and h in a) for a in ALIASES[f]):
                    got[f] = j; used.add(j); break
    return got


def build(rows, header_idx, mapping):
    """Turn raw rows into sections. mapping: field -> column index (None/absent = not mapped).
    Returns (sections, report). Rows that cannot be placed are listed in report['skipped'], never dropped silently."""
    mapped = {f: j for f, j in mapping.items() if j is not None}
    if not any(f in mapped for f in FIELDS):
        raise ImportErr("Map at least one column (Audit Area, Requirements, Procedure / Test or Control) before importing.")
    data = rows[header_idx + 1:] if header_idx is not None else rows
    headers = rows[header_idx] if header_idx is not None else [f"Column {i + 1}" for i in range(len(rows[0]))]
    used_cols = set(mapped.values())
    ignored = [h or f"Column {i + 1}" for i, h in enumerate(headers) if i not in used_cols]
    sections, cur, skipped, heading_rows = [], None, [], 0

    def section(name):
        for s in sections:
            if s["heading"].lower() == name.lower():
                return s
        s = {"heading": name, "rows": []}; sections.append(s); return s

    for n, r in enumerate(data, 1 + (header_idx + 1 if header_idx is not None else 0)):
        vals = {f: r[j] for f, j in mapped.items()}
        nonempty = [c for c in r if c]
        if "section" in mapped and vals.get("section"):
            cur = section(vals["section"])
        elif "section" not in mapped and nonempty and len(set(nonempty)) == 1 and (len(nonempty) >= 2 or
                                                      r.index(nonempty[0]) in {0, mapped.get("area")}):
            # one text only (in the first or Audit Area column), or the same text across merged cells = heading row
            cur = section(nonempty[0]); heading_rows += 1; continue
        if not any(vals.get(f) for f in FIELDS):
            skipped.append((n, " | ".join(nonempty)))
            continue
        if cur is None:
            cur = section("Imported Audit Program")
        cur["rows"].append({f: vals.get(f, "") for f in FIELDS})
    return sections, {"ignored_columns": ignored, "skipped": skipped, "heading_rows": heading_rows,
                      "row_count": sum(len(s["rows"]) for s in sections)}


def merge_into(existing, imported):
    """Append imported sections to a draft's sections; same-named headings (case-insensitive) get the rows appended."""
    out = [{**s, "rows": list(s["rows"])} for s in existing]
    for s in imported:
        for t in out:
            if t["heading"].lower() == s["heading"].lower():
                t["rows"] += [dict(r) for r in s["rows"]]; break
        else:
            out.append({"id": None, "heading": s["heading"], "rows": [dict(r) | {"id": None} for r in s["rows"]]})
    return out
