"""Module J: Word/Excel exports for the new modules. Reuses the styling helpers of export_word/export_excel
(KSG green/white/gold/maroon). Output is editable text/tables - never screenshots."""
import os, re
from io import BytesIO

from docx import Document
from docx.shared import Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import export_excel as xl
import export_word as xw
import theme as th


def safe_name(*parts, ext=""):
    """Sensible file name, e.g. T-1_WP-001_working-paper.docx"""
    s = "_".join(re.sub(r"[^A-Za-z0-9._-]+", "-", str(p)).strip("-") for p in parts if p)
    return f"{s}.{ext}" if ext else s


def _doc(title, label):
    d = Document()
    xw._style_doc(d, label)
    if th.logo_available():
        try: d.add_picture(str(th.LOGO_PATH), height=Cm(2.2))
        except Exception: pass
    d.add_heading(os.getenv("ORG_NAME", "Kenya School of Government"), 0)
    d.add_heading(title, 1)
    xw._rule(d.paragraphs[-1], th.GOLD, "16")
    return d


def _bytes(d):
    b = BytesIO(); d.save(b); return b.getvalue()


def _sign(d, rows):
    xw._section(d, "Sign-off")
    xw._kv(d, rows)


# ------------------------------ working paper ------------------------------
WP_SECTIONS = [("Audit objective", "audit_objective"), ("Control objective", "control_objective"), ("Risk", "risk"),
               ("Audit criteria", "criteria"), ("Procedure / test performed", "procedure"), ("Population", "population"),
               ("Sample selection", "sample_selection"), ("Documents and evidence required", "documents_required"),
               ("Work performed", "work_performed"), ("Evidence references", "evidence_refs"), ("Test results", "test_results"),
               ("Exceptions / deviations (for investigation)", "exceptions"), ("Auditor analysis", "analysis"),
               ("Conclusion", "conclusion")]


def working_paper_docx(wp, ann):
    d = _doc("Audit Working Paper", "Audit Working Paper")
    xw._kv(d, [("Engagement", f"{ann['ref']} - {ann['title']}"), ("Financial year", ann["financial_year"]),
               ("Working paper ref.", wp["wp_ref"]), ("Audit area", wp["area"]), ("AWP procedure ref.", wp["awp_wp_ref"]),
               ("Status", wp["status"]), ("Version", f"v{wp['version']}")])
    for h, k in WP_SECTIONS:
        xw._section(d, h); d.add_paragraph(wp[k] or "[Not yet completed]")
    xw._section(d, "Sample-level testing")
    if wp["samples"]:
        t = d.add_table(rows=1, cols=5); t.style = "Table Grid"; xw._table_borders(t)
        for c, h in zip(t.rows[0].cells, ("Item", "Description", "Evidence ref.", "Result", "Exception note")):
            c.text = h; xw._shade(c, th.GREEN)
            for r in c.paragraphs[0].runs: r.bold = True; r.font.color.rgb = xw._rgb(th.WHITE)
        for s in wp["samples"]:
            for c, v in zip(t.add_row().cells, (s["item_ref"], s["description"], s["evidence_ref"], s["result"], s["exception_note"])): c.text = v or ""
    else:
        d.add_paragraph("[No sample items recorded]")
    _sign(d, [("Prepared by", f"{wp['prepared_by']}   Date: {wp['prepared_date'] or '________'}   Signature: __________"),
              ("Reviewed by", f"{wp['reviewer'] or '__________'}   Date: {wp['review_date'] or '________'}   Signature: __________"),
              ("Reviewer comments", wp["reviewer_comments"])])
    return _bytes(d)


def working_paper_xlsx(wp, ann):
    wb = Workbook(); ws = wb.active; ws.title = "Working Paper"
    rows = [("Engagement", f"{ann['ref']} - {ann['title']}"), ("Financial year", ann["financial_year"]), ("Working paper ref.", wp["wp_ref"]),
            ("Audit area", wp["area"]), ("Status", wp["status"]), ("Version", wp["version"])] + [(h, wp[k]) for h, k in WP_SECTIONS] + \
           [("Prepared by", wp["prepared_by"]), ("Prepared date", wp["prepared_date"]), ("Reviewer", wp["reviewer"]),
            ("Review date", wp["review_date"]), ("Reviewer comments", wp["reviewer_comments"])]
    for i, (k, v) in enumerate(rows, 1):
        a = ws.cell(i, 1, k); a.font = Font(bold=True, color=xl._hx(th.DARK_GREEN)); a.fill = xl._fill(th.LIGHT_GREEN); a.border = xl._B
        b = ws.cell(i, 2, v); b.alignment = Alignment(wrap_text=True, vertical="top"); b.border = xl._B
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 36, 100
    xl._page(ws, landscape=False, label="Audit Working Paper")
    s = wb.create_sheet("Sample Testing")
    xl._head(s, 1, ["Item", "Description", "Evidence ref.", "Result", "Exception note"])
    for r, x in enumerate(wp["samples"], 2):
        for i, v in enumerate((x["item_ref"], x["description"], x["evidence_ref"], x["result"], x["exception_note"]), 1):
            c = s.cell(r, i, v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = xl._B
    for i, w in enumerate((16, 40, 18, 26, 50), 1): s.column_dimensions[get_column_letter(i)].width = w
    xl._page(s, label="Audit Working Paper")
    b = BytesIO(); wb.save(b); return b.getvalue()


# ------------------------------ follow-up ------------------------------
FU_ROWS = [("Original finding / recommendation ref.", "finding_ref"), ("Original recommendation", "original_recommendation"),
           ("Agreed management action", "agreed_action"), ("Action owner / department", "action_owner"),
           ("Original agreed implementation date", "original_due_date"),
           ("MANAGEMENT-REPORTED status", "mgmt_status"), ("Date of management update", "mgmt_update_date"),
           ("Supporting evidence supplied by management", "mgmt_evidence"), ("Auditor verification procedures and comments", "auditor_procedures"),
           ("AUDITOR-VERIFIED status", "auditor_status"), ("Follow-up conclusion", "followup_conclusion"), ("Follow-up date", "followup_date"),
           ("Review status", "review_status"), ("Reviewer", "reviewer"), ("Reviewer comments", "reviewer_comments")]


def followup_docx(fu, ann=None):
    d = _doc("Post-Audit Follow-Up", "Post-Audit Follow-Up")
    xw._kv(d, [("Engagement", f"{ann['ref']} - {ann['title']}" if ann else "[Not linked]")] + [(h, fu[k]) for h, k in FU_ROWS])
    _sign(d, [("Prepared by", f"{fu['created_by']}   Signature: __________  Date: ________"),
              ("Reviewed by", f"{fu['reviewer'] or '__________'}   Signature: __________  Date: ________")])
    return _bytes(d)


def followup_xlsx(fus, anns):
    wb = Workbook(); ws = wb.active; ws.title = "Follow-Up"
    heads = ["Engagement"] + [h for h, _ in FU_ROWS]
    xl._head(ws, 1, heads)
    for r, f in enumerate(fus, 2):
        a = anns.get(f["annual_engagement_id"])
        for i, v in enumerate([f"{a['ref']} - {a['title']}" if a else ""] + [f[k] for _, k in FU_ROWS], 1):
            c = ws.cell(r, i, v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = xl._B
    for i in range(1, len(heads) + 1): ws.column_dimensions[get_column_letter(i)].width = 28
    ws.freeze_panes = "B2"; xl._page(ws, fit=False, label="Post-Audit Follow-Up")
    b = BytesIO(); wb.save(b); return b.getvalue()


# ------------------------------ annual plan register ------------------------------
def register_xlsx(engs):
    wb = Workbook(); ws = wb.active; ws.title = "Engagement Register"
    cols = [("financial_year", "Financial year"), ("ref", "Reference"), ("title", "Title"), ("objective", "Objective"), ("scope", "Scope"),
            ("risk_rating", "Risk rating"), ("audit_type", "Type"), ("planned_start", "Planned start"), ("planned_end", "Planned end"),
            ("department", "Department / auditee"), ("team", "Audit team"), ("status", "Status"), ("source_workbook", "Source workbook"),
            ("source_sheet", "Source sheet"), ("source_row", "Source row"), ("is_test_data", "TEST data?")]
    xl._head(ws, 1, [h for _, h in cols])
    for r, e in enumerate(engs, 2):
        for i, (k, _) in enumerate(cols, 1):
            v = ("YES - not an official engagement" if e[k] else "") if k == "is_test_data" else e[k]
            c = ws.cell(r, i, v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = xl._B
    for i in range(1, len(cols) + 1): ws.column_dimensions[get_column_letter(i)].width = 24
    ws.freeze_panes = "C2"; xl._page(ws, fit=False, label="Annual Audit Plan - Engagement Register")
    b = BytesIO(); wb.save(b); return b.getvalue()
