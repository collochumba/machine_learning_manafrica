"""Modules A and B: annual audit plan / engagement register and the reusable audit-areas catalogue.

Annual engagements (table annual_engagements) are kept separate from the reusable audit areas and from
individual AWP documents (the legacy `engagements` table). Nothing here invents official engagements."""
import io, re, zipfile
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

import database as db

CANDIDATE_AREAS = [  # candidate categories only - NOT proof of approved engagements
    "Information Systems and ICT Infrastructure", "Transport, Fuel and Fleet Management",
    "Administration, Asset and Project Management", "Human Resources and Payroll",
    "Training, Research and Consultancy", "Finance and Accounts (incl. Receivables and Payables)",
    "Hospitality Management", "Supply Chain Management"]
ENG_STATUSES = ["Planned", "In Progress", "Completed", "Deferred", "Cancelled"]
RISK_RATINGS = ["Low", "Medium", "High", "Critical"]
FIELDS = ["ref", "title", "objective", "scope", "risk_rating", "audit_type", "planned_start", "planned_end",
          "department", "team"]
REQUIRED = ["ref", "title"]
RECOMMENDED = ["objective", "scope", "department", "planned_start", "planned_end"]
FIELD_LABELS = {"ref": "Engagement reference", "title": "Engagement title", "objective": "Audit objective",
                "scope": "Scope", "risk_rating": "Risk rating", "audit_type": "Nature / type of audit",
                "planned_start": "Planned start", "planned_end": "Planned end",
                "department": "Responsible department / auditee", "team": "Assigned audit team"}
_ALIASES = {"ref": ["ref", "reference", "engagement ref", "engagement reference", "no", "no.", "code"],
            "title": ["title", "engagement", "engagement title", "audit title", "audit", "name", "audit area"],
            "objective": ["objective", "audit objective", "objectives"], "scope": ["scope", "audit scope"],
            "risk_rating": ["risk", "risk rating", "risk level", "rating"],
            "audit_type": ["type", "audit type", "nature", "nature of audit", "type of audit"],
            "planned_start": ["start", "start date", "planned start", "commencement"],
            "planned_end": ["end", "end date", "planned end", "completion"],
            "department": ["department", "auditee", "responsible department", "unit", "directorate", "division"],
            "team": ["team", "audit team", "assigned to", "auditor", "auditors"]}
MAX_UPLOAD = 5 * 1024 * 1024
TEST_SHEET = "TEST DATA - NOT KSG PLAN"


def _t(v):
    return "" if v is None else str(v).strip()


# ---------------- audit areas (Module B) ----------------
def ensure_seed(user=None):
    """Seed the candidate audit-area categories once (flagged as candidates)."""
    with db.connect() as c:
        if c.execute("SELECT 1 FROM audit_areas").fetchone():
            return 0
        for n in CANDIDATE_AREAS:
            c.execute("INSERT INTO audit_areas(name,description,is_candidate,created_by,created_at,updated_at) "
                      "VALUES(?,?,1,?,?,?)", (n, "Candidate category - reconcile with the approved Annual Audit Plan.",
                                              (user or {}).get("username", "system"), db.now(), db.now()))
        return len(CANDIDATE_AREAS)


def list_areas(include_archived=False):
    with db.connect() as c:
        q = "SELECT * FROM audit_areas" + ("" if include_archived else " WHERE status='Active'") + " ORDER BY name"
        return [dict(r) for r in c.execute(q)]


def save_area(data, user, area_id=None):
    db._need(user, "admin")
    name = _t(data.get("name"))
    if not name:
        raise ValueError("Audit area name is required.")
    with db.connect() as c:
        dup = c.execute("SELECT id FROM audit_areas WHERE lower(name)=lower(?)", (name,)).fetchone()
        if dup and dup["id"] != area_id:
            raise ValueError(f"An audit area named '{name}' already exists.")
        if area_id:
            c.execute("UPDATE audit_areas SET name=?,description=?,category=?,status=?,is_candidate=?,updated_at=? WHERE id=?",
                      (name, _t(data.get("description")), _t(data.get("category")), data.get("status") or "Active",
                       int(data.get("is_candidate", 1)), db.now(), area_id))
        else:
            cur = c.execute("INSERT INTO audit_areas(name,description,category,status,is_candidate,created_by,created_at,updated_at)"
                            " VALUES(?,?,?,?,?,?,?,?)", (name, _t(data.get("description")), _t(data.get("category")),
                                                         data.get("status") or "Active", int(data.get("is_candidate", 1)),
                                                         user["username"], db.now(), db.now()))
            area_id = cur.lastrowid
        db._log(c, user, "save_area", name)
        return area_id


def set_engagement_areas(eng_id, area_ids, user):
    db._need(user, "edit")
    with db.connect() as c:
        c.execute("DELETE FROM engagement_areas WHERE annual_engagement_id=?", (eng_id,))
        for a in dict.fromkeys(area_ids):
            c.execute("INSERT INTO engagement_areas(annual_engagement_id,area_id) VALUES(?,?)", (eng_id, a))
        db._log(c, user, "link_areas", f"engagement {eng_id}: {len(set(area_ids))} area(s)")


def get_engagement_areas(eng_id):
    with db.connect() as c:
        return [dict(r) for r in c.execute(
            "SELECT a.* FROM audit_areas a JOIN engagement_areas l ON l.area_id=a.id WHERE l.annual_engagement_id=? ORDER BY a.name", (eng_id,))]


# ---------------- annual engagements (Module A) ----------------
def _norm_date(v):
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    s = _t(v)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        try:
            date.fromisoformat(s); return s
        except ValueError:
            return s
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return s  # periods such as "Q1" or "Jul-Sep" are kept as supplied


def _is_iso(s):
    return bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", s or ""))


def clean_engagement(d):
    out = {f: _t(d.get(f)) for f in FIELDS}
    out["financial_year"] = _t(d.get("financial_year"))
    for f in ("planned_start", "planned_end"):
        out[f] = _norm_date(d.get(f))
    for f in ("financial_year", "ref", "title"):
        if not out[f]:
            raise ValueError(f"'{FIELD_LABELS.get(f, 'Financial year')}' is required.")
    if _is_iso(out["planned_start"]) and _is_iso(out["planned_end"]) and out["planned_end"] < out["planned_start"]:
        raise ValueError("Planned end cannot be before planned start.")
    if out["risk_rating"] and out["risk_rating"].title() in RISK_RATINGS:
        out["risk_rating"] = out["risk_rating"].title()
    return out


def create_engagement(data, user, source=None, is_test=False):
    db._need(user, "edit")
    d = clean_engagement(data)
    src = source or {}
    with db.connect() as c:
        if c.execute("SELECT 1 FROM annual_engagements WHERE financial_year=? AND lower(ref)=lower(?)",
                     (d["financial_year"], d["ref"])).fetchone():
            raise ValueError(f"Engagement '{d['ref']}' already exists for {d['financial_year']}.")
        cols = ["financial_year"] + FIELDS + ["source_workbook", "source_sheet", "source_row", "is_test_data",
                                              "created_by", "created_at", "updated_at"]
        vals = [d["financial_year"]] + [d[f] for f in FIELDS] + [src.get("workbook", ""), src.get("sheet", ""),
                                                                 src.get("row"), int(is_test), user["username"], db.now(), db.now()]
        cur = c.execute(f"INSERT INTO annual_engagements({','.join(cols)}) VALUES({','.join('?' * len(cols))})", vals)
        db._log(c, user, "create_annual_engagement", f"{d['financial_year']} {d['ref']}")
        return cur.lastrowid


def update_engagement(eid, data, user):
    db._need(user, "edit")
    with db.connect() as c:
        cur = c.execute("SELECT * FROM annual_engagements WHERE id=?", (eid,)).fetchone()
        if not cur:
            raise ValueError("Engagement not found.")
        d = clean_engagement({**dict(cur), **data})
        if c.execute("SELECT 1 FROM annual_engagements WHERE financial_year=? AND lower(ref)=lower(?) AND id<>?",
                     (d["financial_year"], d["ref"], eid)).fetchone():
            raise ValueError(f"Engagement '{d['ref']}' already exists for {d['financial_year']}.")
        status = data.get("status", cur["status"])
        if status not in ENG_STATUSES:
            raise ValueError("Unknown engagement status.")
        sets = ["financial_year"] + FIELDS
        c.execute(f"UPDATE annual_engagements SET {','.join(f + '=?' for f in sets)},status=?,updated_at=? WHERE id=?",
                  [d[f] for f in sets] + [status, db.now(), eid])
        db._log(c, user, "update_annual_engagement", f"{d['financial_year']} {d['ref']}")


def get_engagement(eid):
    with db.connect() as c:
        r = c.execute("SELECT * FROM annual_engagements WHERE id=?", (eid,)).fetchone()
        return dict(r) if r else None


def list_engagements(financial_year=None):
    with db.connect() as c:
        q, p = "SELECT * FROM annual_engagements", ()
        if financial_year:
            q, p = q + " WHERE financial_year=?", (financial_year,)
        return [dict(r) for r in c.execute(q + " ORDER BY financial_year DESC, ref", p)]


def financial_years():
    with db.connect() as c:
        return [r[0] for r in c.execute("SELECT DISTINCT financial_year FROM annual_engagements ORDER BY 1 DESC")]


def linked_records(eid):
    """What depends on this engagement (used to block deletion and to show links)."""
    with db.connect() as c:
        n = lambda q: c.execute(q, (eid,)).fetchone()[0]
        return {"awps": n("SELECT COUNT(*) FROM awp_source WHERE annual_engagement_id=?"),
                "working_papers": n("SELECT COUNT(*) FROM working_papers WHERE annual_engagement_id=?"),
                "followups": n("SELECT COUNT(*) FROM followups WHERE annual_engagement_id=?")}


def delete_engagement(eid, user):
    """Administrator only; refused while AWPs, working papers or follow-ups are linked."""
    db._need(user, "admin")
    links = linked_records(eid)
    if any(links.values()):
        raise ValueError("Cannot delete: this engagement has linked records ("
                         + ", ".join(f"{v} {k.replace('_', ' ')}" for k, v in links.items() if v)
                         + "). Set its status to Cancelled instead.")
    with db.connect() as c:
        e = c.execute("SELECT * FROM annual_engagements WHERE id=?", (eid,)).fetchone()
        if not e:
            raise ValueError("Engagement not found.")
        c.execute("DELETE FROM engagement_areas WHERE annual_engagement_id=?", (eid,))
        c.execute("DELETE FROM annual_engagements WHERE id=?", (eid,))
        db._log(c, user, "delete_annual_engagement", f"{e['financial_year']} {e['ref']}")


# ---------------- Excel import: preview -> mapping -> validate -> confirm ----------------
def validate_upload(filename, data):
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Only .xlsx workbooks can be imported.")
    if not data or len(data) > MAX_UPLOAD:
        raise ValueError("The file is empty or larger than 5 MB.")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("The file is not a valid .xlsx workbook.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("The workbook could not be read. It may be corrupt or password-protected.") from None
    return wb


def sheet_names(filename, data):
    wb = validate_upload(filename, data)
    names = wb.sheetnames
    wb.close()
    return names


def read_sheet(filename, data, sheet, header_row=1, max_rows=5000):
    """Return (headers, rows) where rows are (excel_row_number, [cell values])."""
    wb = validate_upload(filename, data)
    if sheet not in wb.sheetnames:
        raise ValueError("Sheet not found.")
    ws, headers, rows = wb[sheet], None, []
    for i, vals in enumerate(ws.iter_rows(values_only=True), 1):
        if i < header_row:
            continue
        if i == header_row:
            headers = [_t(h) for h in vals]
            continue
        if len(rows) >= max_rows:
            raise ValueError(f"The sheet has more than {max_rows} data rows; split the workbook.")
        if any(_t(v) for v in vals):
            rows.append((i, list(vals)))
    wb.close()
    if not headers or not any(headers):
        raise ValueError("No header row found at the chosen row.")
    return headers, rows


def suggest_mapping(headers):
    """Suggest field -> header index. The user can change every mapping before import."""
    low, used, m = [h.lower() for h in headers], set(), {}
    for f in FIELDS:
        for al in _ALIASES[f]:
            hit = next((i for i, h in enumerate(low) if h == al and i not in used), None)
            if hit is not None:
                m[f] = hit; used.add(hit); break
    return m


def build_preview(headers, rows, mapping, financial_year, workbook_name="", sheet="", is_test=False):
    """Validate mapped rows. Status: New / Incomplete / Duplicate / Invalid. Nothing is saved here."""
    existing = {(r["financial_year"], r["ref"].lower()) for r in list_engagements()}
    seen, out = set(), []
    for xl_row, vals in rows:
        rec = {f: (vals[mapping[f]] if mapping.get(f) is not None and mapping[f] < len(vals) else None) for f in FIELDS}
        rec = {f: (_norm_date(v) if f in ("planned_start", "planned_end") else _t(v)) for f, v in rec.items()}
        issues, status = [], "New"
        missing_req = [FIELD_LABELS[f] for f in REQUIRED if not rec[f]]
        if not financial_year.strip():
            missing_req.append("Financial year")
        if missing_req:
            status, issues = "Invalid", [f"Missing required: {', '.join(missing_req)}"]
        else:
            key = (financial_year.strip(), rec["ref"].lower())
            if key in existing:
                status, issues = "Duplicate", ["Already exists in the register (will not be overwritten)"]
            elif key in seen:
                status, issues = "Duplicate", ["Repeated reference within this workbook"]
            seen.add(key)
            if _is_iso(rec["planned_start"]) and _is_iso(rec["planned_end"]) and rec["planned_end"] < rec["planned_start"]:
                status, issues = "Invalid", ["Planned end is before planned start"]
            if status == "New":
                gaps = [FIELD_LABELS[f] for f in RECOMMENDED if not rec[f]]
                if gaps:
                    status, issues = "Incomplete", [f"Missing: {', '.join(gaps)}"]
                if rec["risk_rating"] and rec["risk_rating"].title() not in RISK_RATINGS:
                    issues.append("Risk rating is not Low/Medium/High/Critical - kept exactly as supplied")
        out.append({"row": xl_row, "status": status, "issues": "; ".join(issues), "financial_year": financial_year.strip(),
                    "workbook": workbook_name, "sheet": sheet, "is_test": is_test, **rec})
    return out


def commit_import(preview, selected_rows, user):
    """Insert only the confirmed rows that are New/Incomplete. Duplicates and invalid rows are never written."""
    db._need(user, "edit")
    chosen, added, skipped = set(selected_rows), 0, []
    with db.connect() as c:  # one transaction: all or nothing
        for r in preview:
            if r["row"] not in chosen:
                continue
            if r["status"] not in ("New", "Incomplete"):
                skipped.append((r["row"], r["status"])); continue
            if c.execute("SELECT 1 FROM annual_engagements WHERE financial_year=? AND lower(ref)=lower(?)",
                         (r["financial_year"], r["ref"])).fetchone():
                skipped.append((r["row"], "Duplicate")); continue
            d = clean_engagement(r)
            cols = ["financial_year"] + FIELDS + ["source_workbook", "source_sheet", "source_row", "is_test_data",
                                                  "created_by", "created_at", "updated_at"]
            vals = [d["financial_year"]] + [d[f] for f in FIELDS] + [r["workbook"], r["sheet"], r["row"], int(r["is_test"]),
                                                                     user["username"], db.now(), db.now()]
            c.execute(f"INSERT INTO annual_engagements({','.join(cols)}) VALUES({','.join('?' * len(cols))})", vals)
            added += 1
        db._log(c, user, "import_annual_plan", f"{added} added, {len(skipped)} skipped")
    return {"added": added, "skipped": skipped}


def make_test_workbook():
    """A clearly labelled TEST workbook. It contains no official KSG engagements."""
    wb = Workbook()
    ws = wb.active
    ws.title = TEST_SHEET
    ws.append(["Engagement Ref", "Engagement Title", "Audit Objective", "Scope", "Risk Rating", "Type of Audit",
               "Start Date", "End Date", "Department", "Audit Team"])
    ws.append(["TEST-001", "TEST ONLY - Fictional ICT review", "Fictional objective for testing the importer.",
               "Fictional scope.", "Medium", "Compliance", "2099-07-01", "2099-08-30", "Test Department", "Test Team"])
    ws.append(["TEST-002", "TEST ONLY - Fictional fleet review", "", "", "High", "", "2099-09-01", "2099-10-15", "", ""])
    ws.append(["TEST-001", "TEST ONLY - Duplicate reference row", "x", "x", "Low", "", "", "", "Test Department", ""])
    ws.append(["", "TEST ONLY - Row with no reference", "x", "x", "Low", "", "", "", "", ""])
    ws2 = wb.create_sheet("README")
    ws2.append(["This workbook is TEST DATA created by the application to demonstrate the import. "
                "It is NOT the KSG Annual Audit Plan and contains no official engagements."])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ---------------- Module C: library <-> audit area links and verification ----------------
VERIFY_STATES = ["Unverified", "Verified"]


def set_library_links(lib_id, area_ids, user):
    db._need(user, "edit")
    with db.connect() as c:
        if not c.execute("SELECT 1 FROM library WHERE id=?", (lib_id,)).fetchone():
            raise ValueError("Library item not found.")
        c.execute("DELETE FROM library_area_links WHERE library_id=?", (lib_id,))
        for a in dict.fromkeys(area_ids):
            c.execute("INSERT INTO library_area_links(library_id,area_id) VALUES(?,?)", (lib_id, a))
        db._log(c, user, "link_library_areas", str(lib_id))


def library_links(lib_id):
    with db.connect() as c:
        return [r[0] for r in c.execute("SELECT area_id FROM library_area_links WHERE library_id=?", (lib_id,))]


def library_for_area(area_id):
    with db.connect() as c:
        return [dict(r) for r in c.execute("SELECT l.* FROM library l JOIN library_area_links k ON k.library_id=l.id WHERE k.area_id=? ORDER BY l.kind,l.title", (area_id,))]


def set_library_meta(lib_id, status, notes, expected, user):
    """'Verified' requires a recorded source document and clause - suggestions are never treated as authoritative."""
    db._need(user, "edit")
    if status not in VERIFY_STATES:
        raise ValueError("Unknown verification status.")
    with db.connect() as c:
        it = c.execute("SELECT * FROM library WHERE id=?", (lib_id,)).fetchone()
        if not it:
            raise ValueError("Library item not found.")
        if status == "Verified" and not (it["doc_title"] and it["clause"]):
            raise ValueError("To mark an item Verified, record its source document title and clause first.")
        c.execute("INSERT INTO library_meta(library_id,verification_status,notes,expected_result) VALUES(?,?,?,?) ON CONFLICT(library_id) "
                  "DO UPDATE SET verification_status=excluded.verification_status,notes=excluded.notes,expected_result=excluded.expected_result",
                  (lib_id, status, _t(notes), _t(expected)))
        c.execute("UPDATE library SET verified=? WHERE id=?", (int(status == "Verified"), lib_id))
        db._log(c, user, "library_verification", f"{lib_id}: {status}")


def get_library_meta(lib_id):
    with db.connect() as c:
        r = c.execute("SELECT * FROM library_meta WHERE library_id=?", (lib_id,)).fetchone()
        return dict(r) if r else {"library_id": lib_id, "verification_status": "Unverified", "notes": "", "expected_result": ""}
