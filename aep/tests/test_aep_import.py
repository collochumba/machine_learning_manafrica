import io, sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import pytest
import aep_import as I

PASTE = ("No.\tAudit Objective / Audit Area\tCriteria\tTest Procedure\tExisting Control\n"
         "ICT Governance\t\t\t\t\n1\tICT policy\tApproved policy exists\tReview the policy\tPeriodic review\n"
         "2\tUser access\tAccess approved\t\"Sample users,\nagree to approvals\"\tSegregation of duties\n"
         "Procurement\t\t\t\t\n1\tTenders\tPPADA 2015\tTest tender files\tTender committee\n")


def test_paste_maps_headings_and_keeps_sections():
    rows = I.parse_text(PASTE)
    h = I.guess_header(rows)
    m = I.auto_map(rows[h])
    assert h == 0 and m == {"area": 1, "requirements": 2, "procedure": 3, "control": 4}
    secs, rep = I.build(rows, h, m)
    assert [(s["heading"], len(s["rows"])) for s in secs] == [("ICT Governance", 2), ("Procurement", 1)]
    assert secs[0]["rows"][1]["procedure"] == "Sample users,\nagree to approvals"
    assert rep["ignored_columns"] == ["No."] and rep["heading_rows"] == 2 and rep["skipped"] == []


def test_section_column_and_skipped_rows_reported():
    t = "Main Audit Area\tAudit Area\tRequirements\tProcedure\tControl\nHR\tPayroll\tR\tP\tC\nFinance\tReceipts\tR\tP\tC\n\tNotes only\t\t\t\nHR\tLeave\tR2\tP2\tC2"
    rows = I.parse_text(t); m = I.auto_map(rows[0]); assert m["section"] == 0
    secs, rep = I.build(rows, 0, m)
    assert [(s["heading"], len(s["rows"])) for s in secs] == [("HR", 2), ("Finance", 2)]
    t2 = "Audit Area\tRequirements\tProcedure\tControl\t Remarks\nA\tR\tP\tC\tnote\n\t\t\t\tOrphan remark"
    rows = I.parse_text(t2); secs, rep = I.build(rows, 0, I.auto_map(rows[0]))
    assert rep["ignored_columns"] == ["Remarks"] and len(rep["skipped"]) == 1 and rep["row_count"] == 1  # nothing silently lost


def test_errors_are_explained():
    for bad in ("", "just one line of words", "a\nb\nc"):
        with pytest.raises(I.ImportErr): I.parse_text(bad)
    with pytest.raises(I.ImportErr): I.build([["a", "b"], ["1", "2"]], 0, {})
    with pytest.raises(I.ImportErr): I.parse_xlsx(b"not a workbook", "Sheet")
    with pytest.raises(I.ImportErr): I.docx_tables(b"not a docx")


def test_xlsx_csv_docx_files():
    import openpyxl, docx
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Program"
    ws.append(["Audit Area", "Criteria", "Audit Procedure", "Control Activity"])
    ws.append(["ICT Governance", None, None, None]); ws.append(["Policy", "R1", "P1", "C1"])
    ws2 = wb.create_sheet("Other"); ws2.append(["x", "y"])
    b = io.BytesIO(); wb.save(b)
    assert I.xlsx_sheets(b.getvalue()) == ["Program", "Other"]
    rows = I.parse_xlsx(b.getvalue(), "Program"); secs, _ = I.build(rows, I.guess_header(rows), I.auto_map(rows[0]))
    assert secs[0]["heading"] == "ICT Governance" and secs[0]["rows"][0]["control"] == "C1"
    csvrows = I.parse_csv(b"Audit Area,Requirements,Procedure,Control\nA,R,P,C\n")
    assert csvrows[1] == ["A", "R", "P", "C"]
    d = docx.Document(); d.add_paragraph("Intro")
    t = d.add_table(rows=3, cols=4)
    for i, v in enumerate(["Audit Area", "Requirements", "Tests to Perform", "Controls"]): t.rows[0].cells[i].text = v
    t.rows[1].cells[0].merge(t.rows[1].cells[3]).text = "Supply Chain Management"
    for i, v in enumerate(["Tenders", "PPADA", "Test files", "Committee"]): t.rows[2].cells[i].text = v
    bb = io.BytesIO(); d.save(bb)
    tabs = I.docx_tables(bb.getvalue()); rows = tabs[0]
    secs, rep = I.build(rows, I.guess_header(rows), I.auto_map(rows[0]))
    assert secs[0]["heading"] == "Supply Chain Management" and secs[0]["rows"][0]["procedure"] == "Test files"


def test_merge_into_existing_sections():
    ex = [{"id": 5, "heading": "HR", "rows": [{"id": 1, "area": "a", "requirements": "", "procedure": "", "control": ""}]}]
    out = I.merge_into(ex, [{"heading": "hr", "rows": [{"area": "b", "requirements": "", "procedure": "", "control": ""}]},
                            {"heading": "New", "rows": []}])
    assert [len(s["rows"]) for s in out] == [2, 0] and out[0]["id"] == 5 and out[1]["id"] is None
