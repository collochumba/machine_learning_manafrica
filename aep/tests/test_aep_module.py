"""Tests for the hardened Audit Engagement Programs module. Google Drive is ALWAYS mocked (FakeDrive): no credentials,
no network, no real Drive documents are touched."""
import io
import sqlite3
import sys
import pathlib
import threading

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import pytest
from conftest_helpers import ADM, AUD, REV, d, use_tmp_db

import aep
import aep_export as ex
import aep_exports as ae
import aep_import as imp
import drive as gd

FOLDER = "FOLDER123"


class FakeDrive:
    """Same interface as drive.DriveClient, with in-memory files and injectable failures."""
    def __init__(self, fail=None, lose_reply=False):
        self.folder_id, self.files, self.fail, self.lose_reply = FOLDER, {}, fail, lose_reply
        self.uploads = self.deletes = 0
        self.n = 0

    def _maybe(self, op):
        if self.fail and (self.fail[0] == op or self.fail[0] == "*"):
            raise gd.DriveError(self.fail[1])

    def check(self):
        self._maybe("check"); return "KSG AEP"

    def find_existing(self, uid, version, fmt):
        self._maybe("find")
        for i, f in self.files.items():
            p = f["props"]
            if p.get("aep_uid") == uid and p.get("aep_version") == str(version) and p.get("fmt") == fmt:
                return {"id": i, "name": f["name"]}

    def list_by_uid(self, uid):
        return [{"id": i, "name": f["name"]} for i, f in self.files.items() if f["props"].get("aep_uid") == uid]

    def upload(self, name, mime, data, props):
        self._maybe("upload")
        self.n += 1; fid = f"F{self.n}"
        self.files[fid] = {"name": name, "props": dict(props), "parents": [FOLDER], "data": data}
        self.uploads += 1
        if self.lose_reply:
            raise gd.DriveError("network")  # uploaded remotely, reply lost
        return fid

    def get_meta(self, fid):
        self._maybe("get")
        if fid not in self.files:
            raise gd.DriveError("not_found")
        f = self.files[fid]
        return {"id": fid, "parents": f["parents"], "appProperties": f["props"]}

    def delete(self, fid):
        self._maybe("delete")
        del self.files[fid]; self.deletes += 1


@pytest.fixture(autouse=True)
def tmpdb(tmp_path, monkeypatch):
    use_tmp_db(tmp_path, monkeypatch)


def meta(**kw):
    return {"title": "Review of ICT Policies", "financial_year": "2026/2027", "campus": "Nairobi", "audit_period": "FY", "status": "Draft"} | kw


SECS = [{"heading": "Information & Communication Technology", "rows": [
    {"area": "Governance", "requirements": "R1", "procedure": "P1", "control": "C1"},
    {"area": "Access", "requirements": "R2", "procedure": "P2", "control": "C2"}]},
    {"heading": "HR & Payroll Management", "rows": [{"area": "Payroll", "requirements": "R3", "procedure": "P3", "control": "C3"}]}]


def make(**kw):
    return aep.save_program(meta(**kw), SECS, AUD)


# ---------------------------------------------------------------- database / concurrency
def test_rows_roundtrip_and_persist_after_restart(tmp_path):
    pid = make()
    p = aep.get_program(pid)
    assert [(r["area"], r["requirements"], r["procedure"], r["control"]) for r in p["sections"][0]["rows"]] == [("Governance", "R1", "P1", "C1"), ("Access", "R2", "P2", "C2")]
    d.init_db()  # simulated application restart: same file, migrations are a no-op, data intact
    again = aep.get_program(pid)
    assert again["title"] == p["title"] and [len(s["rows"]) for s in again["sections"]] == [2, 1] and again["uid"] == p["uid"]
    assert d.DB_PATH.is_file()


def test_db_path_is_configurable_and_sqlite_pragmas(tmp_path, monkeypatch):
    monkeypatch.setattr(d, "DB_PATH", tmp_path / "volume" / "persist.db")
    d.init_db()
    assert (tmp_path / "volume" / "persist.db").is_file()
    with d.connect() as c:
        assert c.execute("PRAGMA foreign_keys").fetchone()[0] == 1
        assert c.execute("PRAGMA busy_timeout").fetchone()[0] == d.BUSY_TIMEOUT_MS
        assert c.execute("PRAGMA journal_mode").fetchone()[0].upper() == "WAL"


def test_concurrent_edit_conflict_is_detected_and_nothing_overwritten():
    pid = make()
    v = aep.get_program(pid)["version"]
    aep.save_program(meta(title="Edited by A"), aep.get_program(pid)["sections"], AUD, pid, expected_version=v)
    with pytest.raises(aep.ConflictError) as e:
        aep.save_program(meta(title="Edited by B (stale)"), SECS, AUD, pid, expected_version=v)
    assert "NOT saved" in str(e.value)
    p = aep.get_program(pid)
    assert p["title"] == "Edited by A" and p["version"] == v + 1 and [len(s["rows"]) for s in p["sections"]] == [2, 1]


def test_save_failure_rolls_back_whole_transaction():
    pid = make()
    bad = [{"id": 99999, "heading": "X", "rows": []}]  # section id that does not belong to this program
    with pytest.raises(ValueError):
        aep.save_program(meta(title="Should not stick"), bad, AUD, pid)
    p = aep.get_program(pid)
    assert p["title"] == "Review of ICT Policies" and len(p["sections"]) == 2


def test_database_locked_gives_clear_error_without_data_loss(monkeypatch):
    pid = make()
    monkeypatch.setattr(d, "BUSY_TIMEOUT_MS", 150)
    blocker = sqlite3.connect(d.DB_PATH, timeout=1); blocker.execute("BEGIN IMMEDIATE")
    try:
        with pytest.raises(d.DatabaseBusy) as e:
            aep.save_program(meta(title="Blocked"), SECS, AUD, pid)
        assert "Nothing was lost" in str(e.value)
    finally:
        blocker.rollback(); blocker.close()
    assert aep.get_program(pid)["title"] == "Review of ICT Policies"
    aep.save_program(meta(title="Now fine"), aep.get_program(pid)["sections"], AUD, pid)  # works once the lock is released


def test_two_threads_saving_neither_corrupts():
    pid = make(); v = aep.get_program(pid)["version"]; out = []
    def go(title):
        try:
            aep.save_program(meta(title=title), aep.get_program(pid)["sections"], AUD, pid, expected_version=v); out.append("ok")
        except aep.ConflictError:
            out.append("conflict")
    ts = [threading.Thread(target=go, args=(t,)) for t in ("T1", "T2")]
    [t.start() for t in ts]; [t.join() for t in ts]
    assert sorted(out) == ["conflict", "ok"]


def test_migration_8_preserves_data_and_drops_out_of_scope_columns():
    with d.connect() as c:
        cols = {r[1] for r in c.execute("PRAGMA table_info(aep_programs)")}
        assert {"work_performed", "test_results", "findings", "prepared_by", "reviewed_by"}.isdisjoint(cols) and {"uid", "version"} <= cols
    assert make()  # new schema accepts saves; legacy upgrade path is covered in test_aep.py


# ---------------------------------------------------------------- paste / validation / security
def test_paste_template_validates_and_builds_the_five_columns():
    text = "No.\tAudit Area\tRequirements\tProcedure / Test\tControl\n1\tAccess\tPolicy\tReview logs\tMFA\n2\tBackup\tDR plan\tTest restore\tOffsite\n"
    rows = imp.parse_text(text)
    h = imp.guess_header(rows)
    secs, rep = imp.build(rows, h, imp.auto_map(rows[h]))
    assert rep["row_count"] == 2 and secs[0]["rows"][0] == {"area": "Access", "requirements": "Policy", "procedure": "Review logs", "control": "MFA"}
    with pytest.raises(imp.ImportErr):
        imp.parse_text("   ")


def test_invalid_inputs_and_permissions():
    for bad in (meta(title=""), meta(financial_year=""), meta(status="Hacked")):
        with pytest.raises(ValueError): aep.save_program(bad, [], AUD)
    with pytest.raises(PermissionError): aep.save_program(meta(), [], REV)
    pid = make()
    with pytest.raises(PermissionError): aep.delete_program(pid, AUD, confirmed=True)  # only Administrator may delete
    with pytest.raises(PermissionError): ae.delete_drive_files(1, AUD, confirmed=True, client=FakeDrive())
    assert aep.get_program(pid)


def test_sql_injection_text_is_stored_literally():
    nasty = "x'); DROP TABLE aep_programs;--"
    pid = aep.save_program(meta(title=nasty), [{"heading": nasty, "rows": [{"area": nasty}]}], AUD)
    p = aep.get_program(pid)
    assert p["title"] == nasty and p["sections"][0]["rows"][0]["area"] == nasty and len(aep.list_programs()) == 1


def test_no_shared_default_password_and_first_admin_flow(monkeypatch):
    monkeypatch.delenv("AWP_DEFAULT_PASSWORD", raising=False)
    with d.connect() as c: c.execute("DELETE FROM users")
    assert d.needs_setup()
    assert d.authenticate("admin", d.LEGACY_DEFAULT_PASSWORD) is None
    with pytest.raises(ValueError): d.create_first_admin("boss", "short")
    with pytest.raises(ValueError): d.create_first_admin("boss", d.LEGACY_DEFAULT_PASSWORD)
    d.create_first_admin("boss", "a-long-unique-pass")
    assert not d.needs_setup() and d.authenticate("boss", "a-long-unique-pass")["role"] == "Administrator"
    with pytest.raises(PermissionError): d.create_first_admin("intruder", "another-long-pass-1")  # not a back door


def test_login_lockout_after_repeated_failures():
    d.save_user("u1", "long-password-12", "Auditor", ADM)
    d._FAILS.clear()
    for _ in range(5): assert d.authenticate("u1", "wrong") is None
    with pytest.raises(PermissionError): d.authenticate("u1", "long-password-12")
    d._FAILS.clear()


def test_legacy_default_password_account_is_flagged_for_forced_change():
    s = "ab" * 16
    with d.connect() as c:
        c.execute("INSERT INTO users(username,salt,pw_hash,role) VALUES('old',?,?,'Auditor')", (s, d._hash(d.LEGACY_DEFAULT_PASSWORD, s)))
    assert d.authenticate("old", d.LEGACY_DEFAULT_PASSWORD)["must_change"] is True


def test_backups_are_pruned(monkeypatch):
    monkeypatch.setattr(d, "BACKUP_KEEP", 3)
    import time
    for i in range(6):
        d.backup_db(f"t{i}"); time.sleep(1.05)
    assert len(list(d.BACKUP_DIR.glob("awp_*.db"))) == 3


# ---------------------------------------------------------------- exports
@pytest.mark.parametrize("fmt", list(ex.FORMATS))
def test_each_format_generates_and_has_columns_in_order_and_no_out_of_scope_fields(fmt):
    p = aep.get_program(make())
    data, name, mime = ex.build(p, fmt)
    assert data and name.startswith("AEP-2026-001_") and name.endswith("." + fmt) and mime == ex.FORMATS[fmt][2]
    if fmt == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        tables = [t for t in doc.tables if t.rows[0].cells[0].text == "No."]
        assert len(tables) == 2 and [c.text for c in tables[0].rows[0].cells] == ex.COLUMNS
        assert [c.text for c in tables[0].rows[1].cells] == ["1", "Governance", "R1", "P1", "C1"]
        text = " ".join(par.text for par in doc.paragraphs) + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    elif fmt == "xlsx":
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        rows = [r for r in ws.iter_rows(values_only=True)]
        hdr = [i for i, r in enumerate(rows) if r[0] == "No."]
        assert len(hdr) == 2 and list(rows[hdr[0]]) == ex.COLUMNS and list(rows[hdr[0] + 1]) == [1, "Governance", "R1", "P1", "C1"]
        text = str(rows)
    elif fmt == "pdf":
        from pypdf import PdfReader
        text = " ".join(pg.extract_text() for pg in PdfReader(io.BytesIO(data)).pages)
        assert text.index("No.") < text.index("Audit Area") < text.index("Requirements") < text.index("Procedure / Test") < text.index("Control")
    else:
        text = data.decode()
        assert text.index("<th style=\"width:5%\">No.</th>") < text.index("<th>Audit Area</th>") < text.index("<th>Procedure / Test</th>") < text.index("<th>Control</th>")
        assert "<!doctype html>" in text
    assert "AUDIT ENGAGEMENT PROGRAMS" in text.upper() and "Review of ICT Policies" in text
    for banned in ("Prepared by", "Prepared By", "Work Performed", "Test Results", "Findings", "Evidence Collected"):
        assert banned not in text


def test_html_escapes_markup_and_excel_neutralises_formulas():
    pid = aep.save_program(meta(title="<script>alert(1)</script>"), [{"heading": "H", "rows": [{"area": "=HYPERLINK(\"http://x\")", "procedure": "<b>p</b>"}]}], AUD)
    p = aep.get_program(pid)
    html_ = ex.build_html(p).decode()
    assert "<script>alert(1)</script>" not in html_ and "&lt;script&gt;" in html_
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(ex.build_xlsx(p))).active
    cell = next(c for row in ws.iter_rows() for c in row if c.value == '=HYPERLINK("http://x")')
    assert cell.data_type == "s"


def test_empty_program_exports_for_every_format():
    p = aep.get_program(aep.save_program(meta(), [], AUD))
    for f in ex.FORMATS: assert ex.build(p, f)[0]


def test_export_failure_creates_no_record_and_keeps_program(monkeypatch):
    pid = make()
    monkeypatch.setitem(ex.BUILDERS, "pdf", lambda p: (_ for _ in ()).throw(RuntimeError("boom")))
    with pytest.raises(ValueError) as e: ae.export_program(pid, "pdf", AUD, FakeDrive())
    assert "could not be generated" in str(e.value) and "boom" not in str(e.value)
    assert ae.list_exports(aep.get_program(pid)["uid"]) == [] and aep.get_program(pid)
    with pytest.raises(ValueError): ae.export_program(pid, "exe", AUD, FakeDrive())


# ---------------------------------------------------------------- Drive upload
def test_upload_success_records_drive_id_and_status():
    pid, fake = make(), FakeDrive()
    r = ae.export_program(pid, "docx", AUD, fake)
    assert r["drive_status"] == "uploaded" and r["data"] and fake.uploads == 1
    row = ae.list_exports(aep.get_program(pid)["uid"])[0]
    assert row["drive_status"] == "uploaded" and row["drive_file_id"] in fake.files and row["filename"] == r["filename"] and row["uploaded_at"]
    f = fake.files[row["drive_file_id"]]
    assert f["props"]["aep_ref"] == "AEP-2026-001" and f["parents"] == [FOLDER]


@pytest.mark.parametrize("kind", ["auth", "permission", "quota", "not_found", "network", "other"])
def test_upload_failures_keep_download_and_never_claim_success(kind):
    pid = make()
    r = ae.export_program(pid, "xlsx", AUD, FakeDrive(fail=("upload", kind)))
    assert r["data"] and r["drive_status"] == "failed" and "Not uploaded" in r["message"] and "Uploaded to Google Drive" not in r["message"]
    row = ae.list_exports(aep.get_program(pid)["uid"])[0]
    assert row["drive_status"] == "failed" and row["drive_error"] == kind and row["drive_file_id"] is None


def test_retry_succeeds_later_and_does_not_duplicate():
    pid, fake = make(), FakeDrive(fail=("upload", "network"))
    ae.export_program(pid, "pdf", AUD, fake)
    fake.fail = None
    assert ae.retry_failed(AUD, fake) == {"uploaded": 1, "failed": 0} and fake.uploads == 1
    assert ae.retry_failed(AUD, fake) == {"uploaded": 0, "failed": 0}  # nothing left to retry
    r2 = ae.export_program(pid, "pdf", AUD, fake)  # unchanged content exported again
    assert fake.uploads == 1 and r2["drive_status"] == "uploaded" and len(ae.list_exports(aep.get_program(pid)["uid"])) == 1


def test_lost_reply_does_not_create_a_duplicate_in_drive():
    pid, fake = make(), FakeDrive(lose_reply=True)
    r = ae.export_program(pid, "html", AUD, fake)  # Drive stored it, but we never saw the reply
    assert r["drive_status"] == "failed" and len(fake.files) == 1
    fake.lose_reply = False
    ae.retry_failed(AUD, fake)
    assert len(fake.files) == 1 and fake.uploads == 1  # retry adopted the existing Drive file
    assert ae.list_exports(aep.get_program(pid)["uid"])[0]["drive_status"] == "uploaded"


def test_changed_program_uploads_a_new_version():
    pid, fake = make(), FakeDrive()
    ae.export_program(pid, "html", AUD, fake)
    p = aep.get_program(pid)
    aep.save_program(meta(title="Changed"), p["sections"], AUD, pid, expected_version=p["version"])
    ae.export_program(pid, "html", AUD, fake)
    assert fake.uploads == 2 and len({f["name"] for f in fake.files.values()}) == 2


def test_missing_credentials_falls_back_to_manual_upload(monkeypatch):
    for k in ("AEP_DRIVE_OAUTH_TOKEN_JSON", "AEP_DRIVE_OAUTH_TOKEN_FILE", "AEP_DRIVE_SERVICE_ACCOUNT_JSON", "AEP_DRIVE_SERVICE_ACCOUNT_FILE"):
        monkeypatch.delenv(k, raising=False)
    assert gd.configured_client() is None and gd.status()[0] == "manual"
    monkeypatch.setenv("AEP_DRIVE_OAUTH_TOKEN_JSON", "{not json")  # corrupt secret: still no crash
    assert gd.configured_client() is None
    pid = make()
    for f in ex.FORMATS:  # all four formats still work and are downloadable
        r = ae.export_program(pid, f, AUD, client=None)
        assert r["data"] and r["drive_status"] == "manual" and "manually" in r["message"]
    assert all(h["drive_status"] == "manual" and h["drive_file_id"] is None for h in ae.list_exports(aep.get_program(pid)["uid"]))
    monkeypatch.setenv("AEP_DRIVE_AUTH", "none")
    assert gd.configured_client() is None


def test_error_messages_never_leak_secrets():
    class Boom(Exception):
        resp = type("R", (), {"status": 401})()
        content = b'{"error":{"errors":[{"reason":"authError"}]}}'
    e = gd.classify(Boom("Bearer ya29.SECRET-TOKEN refresh_token=1//SECRET"))
    assert e.kind == "auth" and "SECRET" not in str(e) and "SECRET" not in e.detail
    assert gd.classify(ConnectionError("x")).kind == "network"


def test_drive_status_check_reports_folder_problems():
    assert gd.status(FakeDrive())[0] == "ready"
    st, msg = gd.status(FakeDrive(fail=("check", "not_found")))
    assert st == "error" and "not found" in msg.lower()


# ---------------------------------------------------------------- deletion
def _two_programs_with_drive_files():
    fake = FakeDrive()
    a, b = make(), aep.save_program(meta(title="Other program"), SECS, AUD)
    ae.export_program(a, "docx", AUD, fake); ae.export_program(a, "pdf", AUD, fake); ae.export_program(b, "docx", AUD, fake)
    return a, b, fake


def test_cancel_deletes_nothing_and_unconfirmed_call_is_refused():
    a, b, fake = _two_programs_with_drive_files()
    with pytest.raises(ValueError): aep.delete_program(a, ADM, confirmed=False)
    with pytest.raises(ValueError): aep.delete_program(a, ADM)  # default is NOT confirmed
    assert len(aep.list_programs()) == 2 and len(fake.files) == 3 and aep.get_program(a)["sections"]


def test_step_a_deletes_db_rows_only_and_writes_manifest():
    a, b, fake = _two_programs_with_drive_files()
    uid_a = aep.get_program(a)["uid"]
    res = aep.delete_program(a, ADM, confirmed=True)
    assert res["program_ref"] == "AEP-2026-001" and len(res["drive_files"]) == 2
    with pytest.raises(ValueError): aep.get_program(a)
    with d.connect() as c:  # cascade removed this program's rows only
        assert c.execute("SELECT COUNT(*) FROM aep_sections").fetchone()[0] == 2
        assert c.execute("SELECT COUNT(*) FROM aep_procedures").fetchone()[0] == 3
        assert c.execute("SELECT COUNT(*) FROM aep_exports WHERE program_uid=?", (uid_a,)).fetchone()[0] == 0
    assert aep.get_program(b)["title"] == "Other program"
    assert len(fake.files) == 3 and fake.deletes == 0  # Drive untouched by step A
    m = ae.get_deletion(res["deletion_id"])
    assert m["program_ref"] == "AEP-2026-001" and m["drive_decision"] == "pending" and {f["status"] for f in m["files"]} == {"retained"}
    assert "R1" not in str(m) and "Governance" not in str(m)  # manifest holds no audit content
    assert [x["id"] for x in ae.open_deletions()] == [res["deletion_id"]]
    assert aep.next_ref("2026/2027") not in ("AEP-2026-001", "AEP-2026-002")  # deleted references are not reused


def test_declining_drive_deletion_leaves_files_in_drive():
    a, b, fake = _two_programs_with_drive_files()
    res = aep.delete_program(a, ADM, confirmed=True)
    ae.decline_drive_deletion(res["deletion_id"], ADM)
    m = ae.get_deletion(res["deletion_id"])
    assert m["drive_decision"] == "declined" and len(fake.files) == 3 and fake.deletes == 0 and {f["status"] for f in m["files"]} == {"retained"}


def test_drive_deletion_needs_its_own_confirmation_and_deletes_only_this_programs_files():
    a, b, fake = _two_programs_with_drive_files()
    res = aep.delete_program(a, ADM, confirmed=True)
    with pytest.raises(ValueError): ae.delete_drive_files(res["deletion_id"], ADM, client=fake)  # not confirmed
    assert len(fake.files) == 3
    out = ae.delete_drive_files(res["deletion_id"], ADM, confirmed=True, client=fake)
    assert [o["status"] for o in out] == ["deleted", "deleted"] and fake.deletes == 2
    assert len(fake.files) == 1 and next(iter(fake.files.values()))["props"]["aep_ref"] == "AEP-2026-002"  # other program's file survives
    assert ae.get_deletion(res["deletion_id"])["drive_decision"] == "confirmed" and ae.open_deletions() == []


def test_drive_deletion_refuses_files_that_do_not_belong_to_the_program():
    a, b, fake = _two_programs_with_drive_files()
    res = aep.delete_program(a, ADM, confirmed=True)
    victim = next(i for i, f in fake.files.items() if f["props"]["aep_ref"] == "AEP-2026-002")
    with d.connect() as c:  # tampered/incorrect manifest entry pointing at another program's file
        c.execute("INSERT INTO aep_deletion_files(deletion_id,drive_file_id,filename,status) VALUES(?,?,?,'retained')", (res["deletion_id"], victim, "other.docx"))
    out = ae.delete_drive_files(res["deletion_id"], ADM, confirmed=True, client=fake)
    assert "skipped_mismatch" in [o["status"] for o in out] and victim in fake.files


def test_drive_deletion_failure_is_recorded_and_can_be_retried():
    a, b, fake = _two_programs_with_drive_files()
    res = aep.delete_program(a, ADM, confirmed=True)
    fake.fail = ("delete", "permission")
    out = ae.delete_drive_files(res["deletion_id"], ADM, confirmed=True, client=fake)
    assert all(o["status"] == "failed" and "permission" in o["detail"].lower() for o in out) and len(fake.files) == 3
    assert ae.get_deletion(res["deletion_id"])["drive_decision"] == "partial" and ae.open_deletions()  # still resolvable
    fake.fail = None
    out = ae.delete_drive_files(res["deletion_id"], ADM, confirmed=True, client=fake)
    assert [o["status"] for o in out] == ["deleted", "deleted"] and len(fake.files) == 1


def test_drive_deletion_handles_missing_file_and_unconfigured_drive():
    a, b, fake = _two_programs_with_drive_files()
    res = aep.delete_program(a, ADM, confirmed=True)
    for fid in [f["drive_file_id"] for f in ae.get_deletion(res["deletion_id"])["files"]]: fake.files.pop(fid)  # already gone / bad id
    out = ae.delete_drive_files(res["deletion_id"], ADM, confirmed=True, client=fake)
    assert {o["status"] for o in out} == {"missing"}
    a2 = make(title="Third")
    ae.export_program(a2, "html", AUD, fake)
    r2 = aep.delete_program(a2, ADM, confirmed=True)
    import os
    os.environ["AEP_DRIVE_AUTH"] = "none"
    try:
        with pytest.raises(ValueError) as e: ae.delete_drive_files(r2["deletion_id"], ADM, confirmed=True)
        assert "not configured" in str(e.value)
    finally:
        os.environ.pop("AEP_DRIVE_AUTH")


def test_drive_deletion_also_finds_tagged_files_whose_upload_reply_was_lost():
    pid, fake = make(), FakeDrive(lose_reply=True)
    ae.export_program(pid, "html", AUD, fake)  # stored in Drive, recorded as failed locally
    res = aep.delete_program(pid, ADM, confirmed=True)
    assert res["drive_files"] == [] and len(fake.files) == 1
    out = ae.delete_drive_files(res["deletion_id"], ADM, confirmed=True, client=fake)
    assert [o["status"] for o in out] == ["deleted"] and fake.files == {}


def test_upgrade_from_v7_database_preserves_programs_sections_rows_and_ids(tmp_path, monkeypatch):
    """A database exactly as the previous release left it (schema v7, with out-of-scope columns populated)."""
    import migrations
    path = tmp_path / "old_v7.db"
    c = sqlite3.connect(path); c.executescript(d.SCHEMA)
    for ver, _, sql in migrations.MIGRATIONS[:7]:
        c.executescript(sql)
    c.execute("INSERT INTO aep_programs(program_ref,title,financial_year,campus,audit_period,main_audit_area,status,work_performed,findings,prepared_by,created_by,created_at) "
              "VALUES('AEP-2025-001','Legacy ICT','2025/2026','Nairobi','FY25','ICT','Completed','old work','old finding','someone','u','2025-07-01')")
    c.execute("INSERT INTO aep_sections(id,program_id,heading,position) VALUES(7,1,'ICT',0),(8,1,'HR',1)")
    c.execute("INSERT INTO aep_procedures(id,section_id,position,area,requirements,procedure,control) VALUES(70,7,0,'A1','R1','P1','C1'),(71,7,1,'A2','R2','P2','C2'),(80,8,0,'A3','R3','P3','C3')")
    c.execute("PRAGMA user_version=7"); c.commit(); c.close()
    monkeypatch.setattr(d, "DB_PATH", path)
    d.init_db()  # backs up first, then applies 8 and 9
    assert list(d.BACKUP_DIR.glob("*pre-migration*.db"))
    p = aep.get_program(1)
    assert (p["program_ref"], p["title"], p["financial_year"], p["campus"], p["status"], p["version"]) == ("AEP-2025-001", "Legacy ICT", "2025/2026", "Nairobi", "Completed", 1) and p["uid"]
    assert [(s["id"], s["heading"], [(r["id"], r["area"], r["control"]) for r in s["rows"]]) for s in p["sections"]] == \
        [(7, "ICT", [(70, "A1", "C1"), (71, "A2", "C2")]), (8, "HR", [(80, "A3", "C3")])]
    assert "work_performed" not in p and "findings" not in p and "prepared_by" not in p
    old = sqlite3.connect(next(d.BACKUP_DIR.glob("*pre-migration*.db")))  # old data still recoverable from the backup
    assert old.execute("SELECT work_performed FROM aep_programs").fetchone()[0] == "old work"
    d.init_db()  # idempotent on restart
    assert aep.save_program(meta(), SECS, AUD) and aep.get_program(2)["program_ref"] == "AEP-2026-001"
