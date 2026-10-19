import os
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
import theme as th
from awp_generator import LABELS
from database import ROW_FIELDS


def _hx(c): return c.lstrip("#")
def _rgb(c): return RGBColor.from_string(_hx(c).upper())


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), _hx(color))
    tcPr.append(shd)


def _table_borders(table, color=th.BORDER):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), _hx(color))
        b.append(e)
    tblPr.append(b)


def _rule(paragraph, color=th.GOLD, size="12"):
    """Bottom border (a coloured rule) under a paragraph or heading."""
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:pBdr")): pPr.remove(old)
    bd = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), size); e.set(qn("w:space"), "1"); e.set(qn("w:color"), _hx(color))
    bd.append(e); pPr.append(bd)


def _style_doc(d, label="Audit Working Plan"):
    st = d.styles
    st["Normal"].font.name, st["Normal"].font.size, st["Normal"].font.color.rgb = "Calibri", Pt(10), _rgb(th.TEXT)
    for name, color, size in (("Title", th.DARK_GREEN, 24), ("Heading 1", th.GREEN, 16), ("Heading 2", th.GREEN, 13), ("Heading 3", th.DARK_GREEN, 11)):
        f = st[name].font
        f.name, f.size, f.bold, f.color.rgb = "Calibri", Pt(size), True, _rgb(color)
        rpr = st[name].element.get_or_add_rPr()
        for rf in rpr.findall(qn("w:rFonts")):  # drop theme-font overrides so Calibri applies
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"): rf.attrib.pop(qn(att), None)
        ppr = st[name].element.get_or_add_pPr()
        for old in ppr.findall(qn("w:pBdr")): ppr.remove(old)  # remove the template's blue rule
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2)
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(f"{th.ORG_NAME} - {label}  |  Page "); r.font.size, r.font.color.rgb = Pt(8), _rgb(th.MUTED)
        r2 = fp.add_run()
        for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if typ:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
            r2._r.append(el)
        r2.font.size, r2.font.color.rgb = Pt(8), _rgb(th.MUTED)


def _kv(doc, pairs):
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    _table_borders(t)
    t.autofit = False
    for k, v in pairs:
        a, b = t.add_row().cells
        a.text, b.text = k, str(v or "")
        _shade(a, th.LIGHT_GREEN)
        a.width, b.width = Cm(5), Cm(12)
        for r in a.paragraphs[0].runs: r.bold = True; r.font.color.rgb = _rgb(th.DARK_GREEN)
        if k in ("Overall risk", "Risk Level") and str(v or "").strip():
            bg, fg, _ = th.risk_style(v); _shade(b, bg)
            for r in b.paragraphs[0].runs: r.bold = True; r.font.color.rgb = _rgb(fg)
        elif k == "Status" and str(v or "").strip():
            bg, fg, _ = th.status_style(v); _shade(b, bg)
            for r in b.paragraphs[0].runs: r.bold = True; r.font.color.rgb = _rgb(fg)
    return t


def _section(d, text):
    h = d.add_heading(text, 2); _rule(h, th.GOLD, "8"); return h


def build_document(eng, rows, comments):
    d = Document()
    _style_doc(d)
    if th.logo_available():
        try: d.add_picture(str(th.LOGO_PATH), height=Cm(2.2))
        except Exception: pass  # an unreadable logo must never block the export
    d.add_heading(os.getenv("ORG_NAME", "Kenya School of Government"), 0)
    d.add_heading("Audit Working Plan", 1)
    _rule(d.paragraphs[-1], th.GOLD, "16")
    _kv(d, [("Reference", eng["ref"]), ("Title", eng["title"]), ("Organisation audited", eng["org"]),
            ("Category", eng["category"]), ("Area / sub-area", f"{eng['area']} / {eng['sub_area']}"),
            ("Audit period", eng["period"]), ("Planned dates", f"{eng['start_date']} to {eng['end_date']}"),
            ("Audit team", eng["team"]), ("Reviewer", eng["reviewer"]), ("Overall risk", eng["risk_level"]),
            ("Version / status", f"v{eng['version']} / {eng['status']}")])
    for h, k in (("Audit Objective", "objective"), ("Scope", "scope"),
                 ("Methodology / Additional Instructions", "instructions"),
                 ("Applicable Policies, Laws and Standards (as entered; unverified)", "policies")):
        _section(d, h); d.add_paragraph(eng[k] or "[Not supplied]")
    _section(d, "Detailed Audit Working Plan")
    for r in rows:
        run = d.add_paragraph().add_run(f"{r['wp_ref']} - {r['area']}"); run.bold = True; run.font.color.rgb = _rgb(th.GREEN)
        d.paragraphs[-1].paragraph_format.keep_with_next = True
        _kv(d, [(LABELS[f], r[f]) for f in ROW_FIELDS[2:]] + [(c, r["extra"].get(c, "")) for c in eng["custom_cols"]])
    _section(d, "Applicable Criteria")
    for c in dict.fromkeys(x["criteria"] for x in rows if x["criteria"]): d.add_paragraph(c, style="List Bullet")
    _section(d, "Evidence Requirements")
    for c in dict.fromkeys(x["evidence_required"] for x in rows if x["evidence_required"]): d.add_paragraph(c, style="List Bullet")
    if comments:
        _section(d, "Review Comments")
        for c in comments: d.add_paragraph(f"{c['created_at']} - {c['author']}: {c['text']}")
    _section(d, "Sign-off")
    _kv(d, [("Prepared by", f"{eng['created_by']}   Signature: __________  Date: ________"),
            ("Reviewed by", f"{eng['reviewer']}   Signature: __________  Date: ________"),
            ("Approval", f"Approved by {eng['approved_by']} on {eng['approved_at']}" if eng["status"] == "Approved" else "NOT APPROVED")])
    buf = BytesIO(); d.save(buf)
    return buf.getvalue()
