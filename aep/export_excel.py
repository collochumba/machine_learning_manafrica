from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from awp_generator import LABELS
from database import ROW_FIELDS
import theme as th

def _hx(c): return c.lstrip("#")  # openpyxl wants hex without '#'
def _fill(c): return PatternFill("solid", fgColor=_hx(c))


_S = Side(style="thin", color=_hx(th.BORDER))
_B = Border(left=_S, right=_S, top=_S, bottom=_S)
_H = _fill(th.GREEN)
_HB = Border(left=_S, right=_S, top=_S, bottom=Side(style="medium", color=_hx(th.GOLD)))  # gold rule under headers
_TXT = _hx(th.TEXT)


def _tag(cell, value, fn):
    """Colour a risk/status cell with the same palette as the dashboard badges."""
    if str(value or "").strip():
        bg, fg, _ = fn(value)
        cell.fill, cell.font = _fill(bg), Font(bold=True, color=_hx(fg))


def _page(ws, landscape=True, fit=True, label="Audit Working Plan"):
    ws.sheet_properties.tabColor = _hx(th.GREEN)
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    if fit:  # narrow sheets only; the 20-column working plan keeps its natural scale so print stays readable
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.center.text = th.ORG_NAME + " - " + label
    ws.oddFooter.center.text = "Page &P of &N"


def _head(ws, r, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v)
        c.font, c.fill, c.border = Font(bold=True, color=_hx(th.WHITE)), _H, _HB
        c.alignment = Alignment(wrap_text=True, vertical="center")


def build_workbook(eng, rows, comments):
    wb = Workbook()
    ws = wb.active; ws.title = "Engagement"
    info = [("Reference", eng["ref"]), ("Title", eng["title"]), ("Organisation", eng["org"]), ("Category", eng["category"]),
            ("Audit area / sub-area", f"{eng['area']} / {eng['sub_area']}"), ("Audit period", eng["period"]),
            ("Planned start", eng["start_date"]), ("Planned end", eng["end_date"]), ("Objective", eng["objective"]),
            ("Scope", eng["scope"]), ("Applicable policies/standards (as entered, unverified)", eng["policies"]),
            ("Audit team", eng["team"]), ("Reviewer/supervisor", eng["reviewer"]), ("Overall risk", eng["risk_level"]),
            ("Additional instructions", eng["instructions"]), ("Version", eng["version"]), ("Status", eng["status"]),
            ("Prepared by", eng["created_by"]),
            ("Approved by", eng["approved_by"] or "NOT APPROVED"), ("Approved at", eng["approved_at"] or "")]
    for i, (k, v) in enumerate(info, 1):
        lab = ws.cell(i, 1, k)
        lab.font, lab.fill = Font(bold=True, color=_hx(th.DARK_GREEN)), _fill(th.LIGHT_GREEN)
        lab.alignment = Alignment(wrap_text=True, vertical="top")
        val = ws.cell(i, 2, v); val.alignment = Alignment(wrap_text=True, vertical="top"); val.font = Font(color=_TXT)
        if k == "Overall risk": _tag(val, v, th.risk_style)
        if k == "Status": _tag(val, v, th.status_style)
        for c in (1, 2): ws.cell(i, c).border = _B
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 38, 90
    _page(ws, landscape=False)

    wp = wb.create_sheet("Working Plan")
    cols = ROW_FIELDS + eng["custom_cols"]
    _head(wp, 1, [LABELS.get(c, c) for c in cols])
    for r, row in enumerate(rows, 2):
        for i, f in enumerate(cols, 1):
            c = wp.cell(r, i, row.get(f, row["extra"].get(f, "")))
            c.alignment, c.border, c.font = Alignment(wrap_text=True, vertical="top"), _B, Font(color=_TXT)
            if f == "risk_level": _tag(c, c.value, th.risk_style)
            elif f == "status": _tag(c, c.value, th.status_style)
    for i, f in enumerate(cols, 1):
        wp.column_dimensions[get_column_letter(i)].width = 14 if f in ("wp_ref", "status", "risk_level") else 30
    wp.freeze_panes = "B2"
    wp.row_dimensions[1].height = 32
    _page(wp, fit=False)
    wp.print_title_rows = "1:1"
    wp.auto_filter.ref = wp.dimensions

    cm = wb.create_sheet("Review & Approval")
    _head(cm, 1, ["Date", "Author", "Comment"])
    for r, c in enumerate(comments, 2):
        for i, v in enumerate((c["created_at"], c["author"], c["text"]), 1):
            x = cm.cell(r, i, v); x.alignment, x.border, x.font = Alignment(wrap_text=True, vertical="top"), _B, Font(color=_TXT)
    cm.column_dimensions["A"].width, cm.column_dimensions["B"].width, cm.column_dimensions["C"].width = 20, 16, 90
    cm.freeze_panes = "A2"
    _page(cm)
    cm.sheet_properties.tabColor = _hx(th.GOLD)
    buf = BytesIO(); wb.save(buf)
    return buf.getvalue()
